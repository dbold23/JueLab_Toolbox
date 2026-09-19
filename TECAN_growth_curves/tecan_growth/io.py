"""
Readers for TECAN Spark plate-reader exports.

Supported inputs
----------------
* **SparkControl / i-control Excel export** (``.xlsx``). The ``Result sheet``
  holds instrument metadata, the plate layout, and the kinetic table whose
  header row starts with ``Cycle Nr.`` followed by ``Time [s]``, ``Temp. [°C]``
  and one column per well (``A1``, ``B1``, … column-major). Lab members often
  copy the table into a second sheet (``Values``) and add a ``Key`` sheet with
  ``Cell, Media, Strain`` columns. All of these are detected by content, never
  by sheet name.
* **Values CSV** — the same kinetic table saved as CSV (``*_VALUES.csv``).
* **Key CSV** — ``Cell, Media, Strain``.

Everything is returned as a :class:`PlateRun` with time in **hours**.
"""
from __future__ import annotations

import io as _io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, BinaryIO, Dict, List, Optional, Sequence, Tuple, Union

import numpy as np
import pandas as pd

from .plate import PlateKey, normalize_well, well_sort_key

PathLike = Union[str, Path]

_CYCLE_RE = re.compile(r"^\s*cycle\s*(nr|no|number|#)?\.?\s*$", re.IGNORECASE)
_TIME_RE = re.compile(r"^\s*time\b", re.IGNORECASE)
_TEMP_RE = re.compile(r"^\s*temp", re.IGNORECASE)

#: Metadata labels harvested from the Result sheet (label -> canonical key).
_META_LABELS = {
    "method name": "method_name",
    "application": "application",
    "device": "device",
    "firmware": "firmware",
    "date": "date",
    "time": "time",
    "system": "system",
    "user": "user",
    "plate": "plate",
    "plate area": "plate_area",
    "part of plate": "part_of_plate",
    "kinetic cycles": "kinetic_cycles",
    "interval time": "interval_time",
    "measurement wavelength": "wavelength_nm",
    "number of flashes": "flashes",
    "target temperature": "target_temperature_c",
    "shaking (orbital) duration": "shaking_duration_s",
    "shaking (orbital) amplitude": "shaking_amplitude_mm",
    "shaking (orbital) frequency": "shaking_frequency_rpm",
    "start time": "start_time",
    "end time": "end_time",
}


@dataclass
class PlateRun:
    """A kinetic OD run: one time axis, one column of raw OD per well."""
    time_h: np.ndarray
    wells: pd.DataFrame                      # index = cycle number (1-based), columns = well ids, values = raw OD
    temperature_c: Optional[np.ndarray] = None
    metadata: Dict[str, Any] = field(default_factory=dict)
    layout: Dict[str, str] = field(default_factory=dict)   # instrument plate-layout labels (e.g. 'A1': 'BL1_1')
    key: Optional[PlateKey] = None
    key_candidates: Dict[str, pd.DataFrame] = field(default_factory=dict)  # sheet name -> raw key table
    source: Dict[str, Any] = field(default_factory=dict)
    warnings: List[str] = field(default_factory=list)

    # ------------------------------------------------------------ properties
    @property
    def n_cycles(self) -> int:
        return len(self.time_h)

    @property
    def well_ids(self) -> List[str]:
        return list(self.wells.columns)

    @property
    def duration_h(self) -> float:
        return float(self.time_h[-1] - self.time_h[0]) if len(self.time_h) else 0.0

    @property
    def interval_h(self) -> float:
        return float(np.median(np.diff(self.time_h))) if len(self.time_h) > 1 else float("nan")

    def summary(self) -> Dict[str, Any]:
        return {
            "n_cycles": self.n_cycles,
            "n_wells": len(self.well_ids),
            "duration_h": round(self.duration_h, 3),
            "interval_min": round(self.interval_h * 60, 2) if self.n_cycles > 1 else None,
            "wavelength_nm": self.metadata.get("wavelength_nm"),
            "date": self.metadata.get("date"),
            "device": self.metadata.get("device"),
            "source": self.source.get("path"),
        }

    def to_values_frame(self) -> pd.DataFrame:
        """Rebuild the TECAN-style values table (Cycle Nr., Time [s], Temp. [°C], wells)."""
        out = pd.DataFrame({"Cycle Nr.": self.wells.index.values, "Time [s]": self.time_h * 3600.0})
        if self.temperature_c is not None:
            out["Temp. [°C]"] = self.temperature_c
        for w in self.well_ids:
            out[w] = self.wells[w].values
        return out


# =============================================================================
# Kinetic table parsing (shared by xlsx and csv)
# =============================================================================

def _is_number(v: Any) -> bool:
    if v is None or isinstance(v, bool):
        return False
    if isinstance(v, (int, float, np.integer, np.floating)):
        return not (isinstance(v, float) and np.isnan(v))
    try:
        float(str(v).strip())
        return True
    except ValueError:
        return False


def _time_to_hours(values: np.ndarray, header: str) -> Tuple[np.ndarray, str]:
    """Convert a time column to hours based on its header unit; returns (hours, unit_found)."""
    h = header.lower()
    if "[s]" in h or "sec" in h or h.strip() in ("time", "time (s)", "time[s]"):
        unit = "s"
    elif "[min]" in h or "min" in h:
        unit = "min"
    elif "[h]" in h or "hour" in h or "hr" in h:
        unit = "h"
    elif "[ms]" in h:
        unit = "ms"
    else:
        unit = "s"
    factor = {"s": 1 / 3600.0, "min": 1 / 60.0, "h": 1.0, "ms": 1 / 3.6e6}[unit]
    return values.astype(float) * factor, unit


def _parse_kinetic_rows(header: Sequence[Any], rows: List[Sequence[Any]], warnings: List[str]) -> Tuple[np.ndarray, pd.DataFrame, Optional[np.ndarray], List[int]]:
    """Turn a header row + data rows into (time_h, wells_df, temperature, cycle_numbers)."""
    header = [("" if h is None else str(h).strip()) for h in header]
    # Locate columns
    cycle_idx = next((i for i, h in enumerate(header) if _CYCLE_RE.match(h)), None)
    time_idx = next((i for i, h in enumerate(header) if _TIME_RE.match(h)), None)
    temp_idx = next((i for i, h in enumerate(header) if _TEMP_RE.match(h)), None)
    if time_idx is None:
        raise ValueError(f"No time column found in kinetic table header: {header[:6]} …")
    well_cols: List[Tuple[int, str]] = []
    for i, h in enumerate(header):
        if i in (cycle_idx, time_idx, temp_idx):
            continue
        w = normalize_well(h)
        if w is not None:
            well_cols.append((i, w))
    if not well_cols:
        raise ValueError("No well columns (A1 … H12) found in kinetic table header")
    dup = pd.Series([w for _, w in well_cols])
    if dup.duplicated().any():
        raise ValueError(f"Duplicate well columns in table: {sorted(set(dup[dup.duplicated()]))}")

    width = max(i for i, _ in well_cols) + 1
    data = np.full((len(rows), len(header)), np.nan, dtype=float)
    cycles: List[int] = []
    kept = 0
    for r in rows:
        r = list(r) + [None] * (len(header) - len(r))
        if not _is_number(r[time_idx]):
            continue
        for j in range(min(len(r), len(header))):
            v = r[j]
            if _is_number(v):
                data[kept, j] = float(str(v).strip()) if not isinstance(v, (int, float, np.integer, np.floating)) else float(v)
        cycles.append(int(float(r[cycle_idx])) if cycle_idx is not None and _is_number(r[cycle_idx]) else kept + 1)
        kept += 1
    data = data[:kept]
    if kept == 0:
        raise ValueError("Kinetic table has no data rows")

    time_h, unit = _time_to_hours(data[:, time_idx], header[time_idx])
    if unit != "s":
        warnings.append(f"Time column '{header[time_idx]}' interpreted as {unit}; converted to hours")
    order = np.argsort(time_h, kind="stable")
    if not np.all(np.diff(time_h) > 0):
        if np.any(np.diff(time_h[order]) <= 0):
            warnings.append("Duplicate time stamps found in kinetic table; rows kept in file order")
        else:
            warnings.append("Time stamps were not increasing; rows re-sorted by time")
            data = data[order]
            time_h = time_h[order]
            cycles = [cycles[i] for i in order]

    wells = pd.DataFrame({w: data[:, i] for i, w in well_cols}, index=pd.Index(cycles, name="cycle"))
    wells = wells[sorted(wells.columns, key=well_sort_key)]
    all_nan = [w for w in wells.columns if wells[w].isna().all()]
    if all_nan:
        warnings.append(f"{len(all_nan)} well column(s) contain no numeric data and were dropped: {', '.join(all_nan[:8])}")
        wells = wells.drop(columns=all_nan)
    temperature = data[:, temp_idx] if temp_idx is not None else None
    return time_h, wells, temperature, cycles


# =============================================================================
# Excel export
# =============================================================================

def _sheet_rows(ws, max_rows: int = 5000) -> List[Tuple[Any, ...]]:
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= max_rows:
            break
        rows.append(tuple(row))
    return rows


def _find_kinetic_header(rows: List[Tuple[Any, ...]]) -> Optional[int]:
    for i, row in enumerate(rows):
        first = row[0] if row else None
        if first is not None and _CYCLE_RE.match(str(first)):
            # must also contain a time column and at least one well
            cells = [("" if c is None else str(c)) for c in row]
            if any(_TIME_RE.match(c) for c in cells) and any(normalize_well(c) for c in cells[2:]):
                return i
    return None


def _kinetic_block(rows: List[Tuple[Any, ...]], header_idx: int) -> List[Tuple[Any, ...]]:
    block = []
    for row in rows[header_idx + 1:]:
        if not row or row[0] is None or not _is_number(row[0]):
            if block:
                break
            continue
        block.append(row)
    return block


def _looks_like_key_header(row: Sequence[Any]) -> bool:
    cells = [str(c).strip().lower() for c in row if c is not None and str(c).strip()]
    if len(cells) < 3:
        return False
    has_well = any(c in ("cell", "well", "well id", "position") for c in cells)
    has_media = any(c in ("media", "medium", "condition", "treatment") for c in cells)
    has_strain = any(c in ("strain", "sample", "isolate", "organism") for c in cells)
    return has_well and has_media and has_strain


def _extract_key_table(rows: List[Tuple[Any, ...]]) -> Optional[pd.DataFrame]:
    for i, row in enumerate(rows[:20]):
        if _looks_like_key_header(row):
            header = [("" if c is None else str(c).strip()) for c in row]
            width = max(j for j, h in enumerate(header) if h) + 1
            body = []
            for r in rows[i + 1:]:
                r = list(r[:width]) + [None] * (width - len(r[:width]))
                if all(c is None or str(c).strip() == "" for c in r):
                    continue
                body.append(r)
            return pd.DataFrame(body, columns=header[:width])
    return None


def _extract_metadata(rows: List[Tuple[Any, ...]], stop_idx: int) -> Tuple[Dict[str, Any], Dict[str, str]]:
    meta, layout = _scan_metadata(rows, 0, stop_idx)
    # The run's real 'End Time' is written after the kinetic table.
    block_len = len(_kinetic_block(rows, stop_idx))
    tail, _ = _scan_metadata(rows, stop_idx + 1 + block_len, len(rows))
    if tail.get("end_time"):
        meta["end_time"] = tail["end_time"]
    return meta, layout


def _scan_metadata(rows: List[Tuple[Any, ...]], start_idx: int, stop_idx: int) -> Tuple[Dict[str, Any], Dict[str, str]]:
    meta: Dict[str, Any] = {}
    layout: Dict[str, str] = {}
    i = start_idx
    while i < stop_idx:
        row = rows[i]
        cells = list(row)
        first = next(((j, c) for j, c in enumerate(cells) if c is not None and str(c).strip() != ""), None)
        if first is None:
            i += 1
            continue
        j, label = first
        label_s = str(label).strip()
        low = label_s.lower().rstrip(":").strip()
        # 'Method name: Growth curve' style (value embedded in the label cell)
        if ":" in label_s and low.split(":")[0].strip() in _META_LABELS and label_s.split(":", 1)[1].strip():
            k, v = label_s.split(":", 1)
            meta[_META_LABELS[k.strip().lower()]] = v.strip()
        elif low in _META_LABELS:
            value = next((c for c in cells[j + 1:] if c is not None and str(c).strip() != ""), None)
            key = _META_LABELS[low]
            if key in ("start_time", "end_time") and key in meta:
                # keep the first start and the last end
                if key == "end_time":
                    meta[key] = _fmt(value)
            else:
                meta[key] = _fmt(value)
        elif label_s == "<>":
            # plate layout grid: header row '<>', 1..12 ; subsequent rows 'A', 'B', ...
            col_numbers = [c for c in cells[j + 1:]]
            k = i + 1
            while k < stop_idx:
                r = list(rows[k])
                rl = r[j] if j < len(r) else None
                if rl is None or not re.match(r"^[A-Pa-p]$", str(rl).strip()):
                    break
                for cnum, val in zip(col_numbers, r[j + 1:]):
                    if cnum is None or val is None:
                        continue
                    try:
                        w = f"{str(rl).strip().upper()}{int(float(cnum))}"
                    except (TypeError, ValueError):
                        continue
                    layout[w] = str(val).strip()
                k += 1
            i = k
            continue
        i += 1
    return meta, layout


def _fmt(v: Any) -> Any:
    if v is None:
        return None
    if hasattr(v, "isoformat"):
        return v.isoformat(sep=" ") if hasattr(v, "hour") else v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def read_tecan_export(path_or_buffer: Union[PathLike, BinaryIO], blank_labels: Optional[Sequence[str]] = None) -> PlateRun:
    """
    Read a TECAN Spark ``.xlsx`` export (any sheet layout) into a :class:`PlateRun`.

    The kinetic table is found by its ``Cycle Nr.`` header on any sheet; when
    several sheets carry a copy, the longest one wins (``Result sheet`` first
    on ties). Sheets whose header row is ``Cell/Media/Strain`` are collected as
    key candidates; the first one is attached as :attr:`PlateRun.key`.
    """
    import openpyxl

    name = str(path_or_buffer) if isinstance(path_or_buffer, (str, Path)) else getattr(path_or_buffer, "name", "<buffer>")
    wb = openpyxl.load_workbook(path_or_buffer, data_only=True, read_only=True)
    warnings: List[str] = []
    candidates: List[Tuple[int, int, str, List[Tuple[Any, ...]], int]] = []  # (n_rows, priority, sheet, rows, header_idx)
    key_tables: Dict[str, pd.DataFrame] = {}
    for ws in wb.worksheets:
        rows = _sheet_rows(ws)
        h = _find_kinetic_header(rows)
        if h is not None:
            n = len(_kinetic_block(rows, h))
            priority = 1 if ws.title.strip().lower().startswith("result") else 0
            candidates.append((n, priority, ws.title, rows, h))
        kt = _extract_key_table(rows)
        if kt is not None and len(kt):
            key_tables[ws.title] = kt
    wb.close()
    if not candidates:
        raise ValueError(
            f"{name}: no kinetic table found. Expected a sheet with a header row starting 'Cycle Nr.', "
            "'Time [s]' and well columns A1 … H12 (the SparkControl 'Result sheet')."
        )
    candidates.sort(key=lambda c: (c[0], c[1]), reverse=True)
    n, _, sheet, rows, h = candidates[0]
    others = [c for c in candidates[1:] if c[0] != n]
    if others:
        warnings.append(
            "Several sheets contain kinetic tables of different length: "
            + ", ".join(f"{c[2]} ({c[0]} cycles)" for c in candidates)
            + f"; using '{sheet}'."
        )
    time_h, wells, temperature, cycles = _parse_kinetic_rows(rows[h], _kinetic_block(rows, h), warnings)
    meta, layout = _extract_metadata(rows, h)
    if not meta and sheet.strip().lower() != "result sheet":
        # metadata lives on the Result sheet even when the values were taken from a copy
        for c in candidates:
            if c[2].strip().lower().startswith("result"):
                meta, layout = _extract_metadata(c[3], c[4])
                break

    key = None
    key_sheet = None
    if key_tables:
        # prefer a sheet literally called key/schematic, else the first found
        ordered = sorted(key_tables, key=lambda s: (0 if "key" in s.lower() else 1 if "schem" in s.lower() else 2, s))
        for s in ordered:
            try:
                key = PlateKey.from_dataframe(key_tables[s], source=f"{name}:{s}", **({"blank_labels": blank_labels} if blank_labels else {}))
                key_sheet = s
                break
            except ValueError as e:
                warnings.append(f"Sheet '{s}' looked like a key but could not be parsed: {e}")
    run = PlateRun(
        time_h=time_h, wells=wells, temperature_c=temperature, metadata=meta, layout=layout,
        key=key, key_candidates=key_tables,
        source={"path": name, "format": "xlsx", "kinetic_sheet": sheet, "key_sheet": key_sheet, "n_key_sheets": len(key_tables)},
        warnings=warnings,
    )
    return run


# =============================================================================
# CSV inputs
# =============================================================================

def _read_text_table(path_or_buffer: Union[PathLike, BinaryIO, str]) -> List[Tuple[Any, ...]]:
    """Read a CSV/TSV as raw rows (strings), tolerant to encodings and delimiters."""
    if isinstance(path_or_buffer, (str, Path)) and Path(str(path_or_buffer)).exists():
        raw = Path(str(path_or_buffer)).read_bytes()
    elif hasattr(path_or_buffer, "read"):
        raw = path_or_buffer.read()
        if isinstance(raw, str):
            raw = raw.encode("utf-8")
    else:
        raw = str(path_or_buffer).encode("utf-8")
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Could not decode text file")
    import csv
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delim = dialect.delimiter
    except csv.Error:
        delim = ","
    reader = csv.reader(_io.StringIO(text), delimiter=delim)
    return [tuple(r) for r in reader]


def read_values_csv(path_or_buffer: Union[PathLike, BinaryIO]) -> PlateRun:
    """Read a values CSV (``Cycle Nr., Time [s], Temp. [°C], A1 … H12``) into a :class:`PlateRun`."""
    name = str(path_or_buffer) if isinstance(path_or_buffer, (str, Path)) else getattr(path_or_buffer, "name", "<buffer>")
    rows = _read_text_table(path_or_buffer)
    h = _find_kinetic_header(rows)
    if h is None:
        # Some exports omit 'Cycle Nr.'; accept a header with a time column and wells
        for i, row in enumerate(rows[:50]):
            cells = [str(c) for c in row]
            if any(_TIME_RE.match(c) for c in cells) and sum(1 for c in cells if normalize_well(c)) >= 1:
                h = i
                break
    if h is None:
        raise ValueError(f"{name}: no kinetic table header (Cycle Nr., Time [s], A1 …) found")
    warnings: List[str] = []
    time_h, wells, temperature, cycles = _parse_kinetic_rows(rows[h], _kinetic_block(rows, h), warnings)
    return PlateRun(time_h=time_h, wells=wells, temperature_c=temperature, metadata={}, layout={},
                    source={"path": name, "format": "csv"}, warnings=warnings)


def read_key(path_or_buffer: Union[PathLike, BinaryIO], sheet: Optional[str] = None, **kwargs) -> PlateKey:
    """Read a plate key from CSV or from a sheet of an xlsx file."""
    name = str(path_or_buffer) if isinstance(path_or_buffer, (str, Path)) else getattr(path_or_buffer, "name", "<buffer>")
    if name.lower().endswith((".xlsx", ".xlsm")) or sheet is not None:
        import openpyxl
        wb = openpyxl.load_workbook(path_or_buffer, data_only=True, read_only=True)
        sheets = [wb[sheet]] if sheet else wb.worksheets
        for ws in sheets:
            kt = _extract_key_table(_sheet_rows(ws, 500))
            if kt is not None:
                wb.close()
                return PlateKey.from_dataframe(kt, source=f"{name}:{ws.title}", **kwargs)
        wb.close()
        raise ValueError(f"{name}: no sheet with Cell/Media/Strain columns found")
    rows = _read_text_table(path_or_buffer)
    kt = _extract_key_table(rows)
    if kt is None:
        raise ValueError(f"{name}: expected columns Cell, Media, Strain")
    return PlateKey.from_dataframe(kt, source=name, **kwargs)


def load_plate(values: Union[PathLike, BinaryIO], key: Optional[Union[PathLike, BinaryIO, PlateKey]] = None) -> PlateRun:
    """
    One-call loader: ``values`` may be an xlsx export or a values CSV; ``key`` an
    optional CSV/xlsx/PlateKey that overrides any key sheet found in the export.
    """
    name = str(values) if isinstance(values, (str, Path)) else getattr(values, "name", "")
    if name.lower().endswith((".xlsx", ".xlsm")):
        run = read_tecan_export(values)
    else:
        run = read_values_csv(values)
    if key is not None:
        run.key = key if isinstance(key, PlateKey) else read_key(key)
        run.source["key_sheet"] = getattr(run.key, "source", None)
    return run
